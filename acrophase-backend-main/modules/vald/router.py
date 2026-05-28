from __future__ import annotations

from collections import OrderedDict
from datetime import date, datetime
from io import BytesIO
import json
import logging
import re
import time
from typing import Optional

from docx import Document
from fastapi import APIRouter, Body, File, Form, Query, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse

from modules.common.config import client
from modules.common.responses import error_response, success_response
from modules.vald.client import ValdApiClient
from modules.vald.repository import ValdRepository
from modules.vald.renderer import render_vald_joint_pdf
from modules.vald.sync import ValdSyncService, pick

router = APIRouter()
logger = logging.getLogger(__name__)
_profile_refresh_at = 0.0
_PROFILE_REFRESH_SECONDS = 300
_athlete_profile_cache: dict[str, dict] = {}


@router.post("/vald/sync")
async def sync_vald_data():
    try:
        result = await ValdSyncService().sync_all()
        return success_response(data=result, message="VALD sync completed")
    except Exception as exc:
        return JSONResponse(
            status_code=500,
            content=error_response(message="VALD sync failed", errors=[str(exc)]),
        )


@router.get("/vald/athletes")
async def search_vald_athletes(q: str = Query("", max_length=100)):
    try:
        await refresh_vald_profiles_if_needed(force=not q.strip())
        athletes = await ValdRepository().lookup_athletes_by_name(q.strip())
        return {"athletes": athletes}
    except Exception as exc:
        logger.warning("VALD athlete search refresh failed: %s", exc)
        athletes = await ValdRepository().lookup_athletes_by_name(q.strip())
        return {"athletes": athletes}


@router.get("/vald/athletes/{athlete_id}/assessment-dates")
async def vald_assessment_dates(
    athlete_id: str, device: Optional[str] = None
):
    device_key = {"dynamometer": "dynamo", "forcedecks": "forcedecks"}.get(
        device or "", device
    )
    dates = await ValdRepository().fetch_assessment_dates(
        athlete_id, device=device_key
    )
    return {"assessment_dates": dates}


@router.get("/athlete-report")
async def athlete_report(
    name: str = Query(..., min_length=1),
    athlete_id: Optional[str] = None,
    date_from: Optional[str] = Query(None, alias="from"),
    date_to: Optional[str] = Query(None, alias="to"),
    test_type: Optional[str] = None,
    latest_only: bool = False,
    device: Optional[str] = None,
    assessment_date: Optional[str] = None,
    sport: Optional[str] = None,
    refresh: bool = False,
):
    if assessment_date:
        try:
            date.fromisoformat(assessment_date)
        except ValueError:
            return JSONResponse(
                status_code=400,
                content=error_response(
                    message="Invalid assessment date",
                    errors=["assessment_date must use YYYY-MM-DD format"],
                ),
            )
    repo = ValdRepository()
    try:
        await refresh_vald_profiles_if_needed()
    except Exception as exc:
        logger.warning("VALD profile refresh failed during report lookup: %s", exc)
    selected_athlete = await repo.fetch_athlete_by_id(athlete_id) if athlete_id else None
    matches = [selected_athlete] if selected_athlete else await repo.lookup_athletes_by_name(name.strip())
    exact_matches = (
        [selected_athlete]
        if selected_athlete
        else [
            athlete
            for athlete in matches
            if athlete.get("name", "").strip().lower() == name.strip().lower()
        ]
    )

    if not exact_matches:
        if matches:
            return JSONResponse(
                status_code=409,
                content={
                    "message": "Multiple or partial athlete matches found",
                    "candidates": matches,
                },
            )
        return JSONResponse(
            status_code=404,
            content=error_response(
                message="Athlete not found",
                errors=[f"No active VALD athlete found for name '{name}'"],
            ),
        )

    if len(exact_matches) > 1:
        return JSONResponse(
            status_code=409,
            content={
                "message": "Multiple athlete matches found",
                "candidates": exact_matches,
            },
        )

    athlete = exact_matches[0]
    device_key = {"dynamometer": "dynamo", "forcedecks": "forcedecks"}.get(
        device or "", device
    )
    tests = await repo.fetch_tests_for_athlete(
        athlete["vald_id"], date_from=date_from, date_to=date_to, test_type=test_type,
        device=device_key, assessment_date=assessment_date,
    )
    stored_metrics = await repo.fetch_metrics_for_tests(
        [test["vald_test_id"] for test in tests]
    )
    must_hydrate = not tests or not any(stored_metrics.values())
    if assessment_date and device_key in {"dynamo", "forcedecks"} and (refresh or must_hydrate):
        try:
            api = ValdApiClient()
            tenants = await api.tenants()
            tenant_id = pick(tenants[0], "tenantId", "id", "teamId") if tenants else None
            if tenant_id:
                service = ValdSyncService()
                await service.refresh_athlete_date(
                    tenant_id, athlete["vald_id"], assessment_date, device_key
                )
        except Exception as exc:
            logger.warning("VALD selected-date refresh failed for %s: %s", athlete["vald_id"], exc)
        tests = await repo.fetch_tests_for_athlete(
            athlete["vald_id"], date_from=date_from, date_to=date_to, test_type=test_type,
            device=device_key, assessment_date=assessment_date,
        )
    if latest_only:
        tests = most_recent_per_test_type(tests)

    metrics_by_test = await repo.fetch_metrics_for_tests(
        [test["vald_test_id"] for test in tests]
    )
    if not any(metrics_by_test.values()) and tests:
        try:
            api = ValdApiClient()
            tenants = await api.tenants()
            tenant_id = pick(tenants[0], "tenantId", "id", "teamId") if tenants else None
            if tenant_id:
                service = ValdSyncService()
                service.repo = repo
                repaired = await service.repair_empty_dynamo_metrics(
                    tenant_id, tests, metrics_by_test
                )
                if repaired:
                    metrics_by_test = await repo.fetch_metrics_for_tests(
                        [test["vald_test_id"] for test in tests]
                    )
        except Exception as exc:
            logger.warning("VALD metric repair failed for %s: %s", athlete["vald_id"], exc)
    try:
        athlete = await enrich_athlete_from_vald_cached(athlete)
    except Exception as exc:
        logger.warning("VALD demographic enrichment failed for %s: %s", athlete["vald_id"], exc)

    response = {
        "athlete": athlete,
        "context": {"assessment_date": assessment_date, "sport": sport},
        "dynamometer": {"tests": [], "joints": []},
        "forcedecks": {"tests": [], "joints": []},
    }

    for test in consolidate_bilateral_tests(tests, metrics_by_test):
        section = "dynamometer" if test["device"] == "dynamo" else "forcedecks"
        metrics = test.get("combined_metrics") or metrics_by_test.get(test["vald_test_id"], [])
        selected_metrics = select_key_metrics(
            test.get("test_type") or "", metrics, device=test.get("device") or ""
        )
        if not selected_metrics:
            continue
        response[section]["tests"].append(
            {
                "vald_test_id": test["vald_test_id"],
                "test_type": test.get("test_type"),
                "test_date": test.get("test_date"),
                "metrics": selected_metrics,
                "available_metric_count": len(metrics),
            }
        )

    for section in ("dynamometer", "forcedecks"):
        response[section]["joints"] = group_tests_by_joint(response[section]["tests"])

    return response


async def refresh_vald_profiles_if_needed(force: bool = False) -> None:
    global _profile_refresh_at
    if not force and time.monotonic() - _profile_refresh_at < _PROFILE_REFRESH_SECONDS:
        return
    service = ValdSyncService()
    tenants = await service.api.tenants()
    tenant_id = pick(tenants[0], "tenantId", "id", "teamId") if tenants else None
    if tenant_id:
        await service.sync_profiles(tenant_id)
        for profile in service.latest_profiles:
            vald_id = pick(profile, "profileId", "athleteId", "id")
            if vald_id:
                _athlete_profile_cache[vald_id] = profile
        _profile_refresh_at = time.monotonic()


async def enrich_athlete_from_vald_cached(athlete: dict) -> dict:
    vald_id = athlete["vald_id"]
    if vald_id in _athlete_profile_cache:
        return enriched_athlete(athlete, _athlete_profile_cache[vald_id])
    api = ValdApiClient()
    tenants = await api.tenants()
    tenant_id = pick(tenants[0], "tenantId", "id", "teamId") if tenants else None
    if not tenant_id:
        return athlete
    profiles = await api.profiles(tenant_id, [vald_id])
    profile = next(
        (item for item in profiles if pick(item, "profileId", "athleteId", "id") == vald_id),
        None,
    )
    if not profile:
        return athlete
    enriched = enriched_athlete(athlete, profile)
    _athlete_profile_cache[vald_id] = profile
    return enriched


@router.post("/vald/joint-interpretations")
async def generate_joint_interpretations(data: dict = Body(...)):
    athlete = data.get("athlete") or {}
    joints = data.get("joints") or []
    report_type = data.get("report_type") or "VALD"
    sport = str(data.get("sport") or "").strip() or "sport not provided"
    assessment_date = data.get("assessment_date") or "not provided"
    interpretations: dict[str, str] = {}

    for joint in joints:
        joint_name = str(joint.get("joint") or "").strip()
        tests = joint.get("tests") or []
        metric_lines = [
            format_metric_for_prompt(test.get("test_type", "Test"), metric)
            for test in tests
            for metric in (test.get("metrics") or [])
        ][:80]
        if not joint_name or not metric_lines:
            continue

        prompt = f"""
You are an elite sports physiotherapist interpreting VALD bilateral testing data for a high-performance report.

Athlete: {athlete.get("name", "Athlete")}
Age: {athlete.get("age_years") or "not available"}
Height: {athlete.get("height_cm") or "not available"} cm
Weight: {athlete.get("weight_kg") or "not available"} kg
Report system: {report_type}
Sport: {sport}
Assessment date: {assessment_date}
Joint/region: {joint_name}

Measured data:
{chr(10).join(metric_lines)}

Write a joint-specific assessment in 170-220 words, with the judgement and tone of a world-class sports physiotherapist.
Use exactly these Markdown headings:
**What Looks Good**
**Main Asymmetry**
**Why It Matters For {sport}**
**Performance Focus**

Mention which movements look good or acceptable from a symmetry/balance perspective when their asymmetry is within the operational 10% band.
For scalar metrics without right-left asymmetry labels, describe them as reported outputs rather than good, poor, high, or low unless a benchmark is supplied.
Mention the highest-asymmetry movement(s), the direction, and why those joint actions matter in {sport}.
It is acceptable to say that improving the flagged movement quality and right-left balance can help improve sport performance, provided you avoid injury diagnosis.
Use only the supplied data. Describe observed left-right differences when shown, and connect them to sport-specific movement tasks.
If you mention a percentage or value, keep it attached to the exact metric name provided in the measured data. Do not move a value from RFD, impulse, force, ROM, stiffness, RSI, or jump height onto another metric.
Do not invent review intervals, return-to-play advice, or exact exercise prescriptions.
Do not call a value significant, deficient, abnormal, risky, or injury-related without a supplied benchmark.
Do not diagnose injury. Frame actions as performance-focused options for coach, therapist, or practitioner review.
Green or red asymmetry screening labels use an operational 10% review threshold, not an age- or sport-specific VALD norm.
Briefly state that absolute strength/ROM quality still needs VALD Hub norms or an approved team benchmark for age-matched interpretation.
"""
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You write high-performance sports physiotherapy assessments. "
                        "Be specific, practical, sport-relevant, and performance-focused without diagnosing injury."
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            temperature=0.1,
        )
        interpretation = constrain_interpretation_language(
            response.choices[0].message.content or ""
        )
        if len(interpretation.split()) > 245:
            interpretation = compact_interpretation(
                interpretation, joint_name, sport
            )
        interpretations[joint_name] = interpretation

    return {"interpretations": interpretations}


@router.post("/vald/final-pdf")
async def generate_vald_final_pdf(data: dict = Body(...)):
    buffer = BytesIO()
    await render_vald_joint_pdf(buffer, data)
    buffer.seek(0)
    athlete_name = re.sub(
        r"[^A-Za-z0-9_-]+", "_", str((data.get("athlete") or {}).get("name") or "athlete")
    ).strip("_")
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename={athlete_name}_VALD_Joint_Report.pdf"
        },
    )


@router.post("/vald/draft-word")
async def generate_vald_draft_word(data: dict = Body(...)):
    doc = Document()
    athlete = data.get("athlete") or {}
    joints = data.get("joints") or []
    interpretations = data.get("interpretations") or {}
    athlete_name = str(athlete.get("name") or "Athlete")

    doc.add_heading("VALD Report Draft", level=1)
    doc.add_paragraph(
        "Edit only the narrative inside each START/END interpretation block, then upload this document back into AcroReports to generate the final PDF."
    )

    profile_table = doc.add_table(rows=0, cols=2)
    profile_table.style = "Table Grid"
    for label, value in (
        ("Athlete", athlete_name),
        ("Sport", data.get("sport") or "-"),
        ("Assessment date", data.get("assessment_date") or "-"),
        ("Report type", data.get("report_type") or "VALD Report"),
        ("Age", athlete.get("age_years") or "-"),
        ("Height", f"{athlete.get('height_cm')} cm" if athlete.get("height_cm") else "-"),
        ("Weight", f"{athlete.get('weight_kg')} kg" if athlete.get("weight_kg") else "-"),
    ):
        cells = profile_table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)

    for joint in joints:
        joint_name = str(joint.get("joint") or "Section")
        doc.add_page_break()
        doc.add_heading(f"{joint_name} Assessment", level=2)
        for test in joint.get("tests") or []:
            doc.add_heading(str(test.get("test_type") or "Test"), level=3)
            table = doc.add_table(rows=1, cols=6)
            table.style = "Table Grid"
            headers = ["Metric", "Value", "Right", "Left", "Unit", "Asymmetry"]
            for cell, header in zip(table.rows[0].cells, headers):
                cell.text = header
            for metric in test.get("metrics") or []:
                cells = table.add_row().cells
                cells[0].text = str(metric.get("name") or "-")
                cells[1].text = format_doc_value(metric.get("value"))
                cells[2].text = format_doc_value(metric.get("right_value"))
                cells[3].text = format_doc_value(metric.get("left_value"))
                cells[4].text = str(metric.get("unit") or "-")
                cells[5].text = compact_asymmetry_label(metric)

        safe_joint = safe_doc_key(joint_name)
        doc.add_paragraph(f"[START_VALD_INTERPRETATION_{safe_joint}]")
        add_markdown_to_doc(doc, interpretations.get(joint_name, ""))
        doc.add_paragraph(f"[END_VALD_INTERPRETATION_{safe_joint}]")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    filename = re.sub(r"[^A-Za-z0-9_-]+", "_", athlete_name).strip("_") or "Athlete"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}_VALD_Draft.docx"},
    )


@router.post("/vald/upload-final-word")
async def parse_vald_final_word(
    file: UploadFile = File(...), joints_json: str = Form(...)
):
    if not file.filename or not file.filename.lower().endswith(".docx"):
        return JSONResponse(
            status_code=400,
            content=error_response(
                message="Invalid file format",
                errors=["Please upload a .docx file."],
            ),
        )

    try:
        joints = json.loads(joints_json)
    except json.JSONDecodeError:
        return JSONResponse(
            status_code=400,
            content=error_response(
                message="Invalid report context",
                errors=["joints_json must be valid JSON."],
            ),
        )

    content = await file.read()
    doc = Document(BytesIO(content))
    section_map = {
        safe_doc_key(str(joint.get("joint") or "Section")): str(joint.get("joint") or "Section")
        for joint in joints
    }
    interpretations = extract_vald_interpretations(doc, section_map)
    return {"interpretations": interpretations}


def most_recent_per_test_type(tests: list[dict]) -> list[dict]:
    selected = OrderedDict()
    for test in tests:
        key = (test.get("device"), test.get("test_type"))
        if key not in selected:
            selected[key] = test
    return list(selected.values())


def enriched_athlete(athlete: dict, profile: dict) -> dict:
    date_of_birth = pick(profile, "dateOfBirth", "date_of_birth")
    return {
        **athlete,
        "date_of_birth": date_of_birth,
        "age_years": calculate_age(date_of_birth),
        "height_cm": pick(profile, "heightCm", "height", "height_cm"),
        "weight_kg": pick(profile, "weightKg", "weight", "weight_kg"),
        "demographics_note": (
            "The VALD External Profiles API supplies date of birth for age calculation; "
            "height and weight are not returned by its current profile retrieval response."
        ),
    }


def calculate_age(date_of_birth: str | None) -> int | None:
    if not date_of_birth:
        return None
    try:
        born = datetime.fromisoformat(date_of_birth.replace("Z", "+00:00")).date()
    except ValueError:
        return None
    today = date.today()
    return today.year - born.year - ((today.month, today.day) < (born.month, born.day))


def consolidate_bilateral_tests(
    tests: list[dict], metrics_by_test: dict[str, list[dict]]
) -> list[dict]:
    combined: OrderedDict[tuple[str, str, str], dict] = OrderedDict()
    for test in tests:
        test_type = str(test.get("test_type") or "")
        base_type = re.sub(r"\s*-\s*(Right|Left)\s*$", "", test_type, flags=re.IGNORECASE)
        date_key = str(test.get("test_date") or "")[:10]
        key = (str(test.get("device") or ""), base_type, date_key)
        if test.get("device") != "dynamo":
            key = (str(test.get("device") or ""), test_type, str(test.get("vald_test_id") or ""))
        if key not in combined:
            combined[key] = {
                **test,
                "test_type": base_type,
                "combined_metrics": [],
                "source_test_ids": [],
            }
        combined[key]["combined_metrics"].extend(
            metrics_by_test.get(test.get("vald_test_id"), [])
        )
        combined[key]["source_test_ids"].append(test.get("vald_test_id"))
    return list(combined.values())


def group_tests_by_joint(tests: list[dict]) -> list[dict]:
    grouped: OrderedDict[str, list[dict]] = OrderedDict()
    for test in tests:
        joint = joint_from_test_type(test.get("test_type") or "")
        grouped.setdefault(joint, []).append(test)
    return [
        {
            "joint": joint,
            "tests": joint_tests,
            "metric_count": sum(len(test.get("metrics") or []) for test in joint_tests),
            "available_metric_count": sum(
                test.get("available_metric_count", len(test.get("metrics") or []))
                for test in joint_tests
            ),
            "last_test_date": next(
                (test.get("test_date") for test in joint_tests if test.get("test_date")),
                None,
            ),
        }
        for joint, joint_tests in grouped.items()
    ]


def select_key_metrics(
    test_type: str, metrics: list[dict], limit: int = 5, device: str = ""
) -> list[dict]:
    if device == "forcedecks":
        return select_forcedecks_metrics(test_type, metrics, limit=limit)

    asymmetry = next(
        (
            metric
            for metric in metrics
            if "asymmetry" in str(metric.get("name") or "").lower()
        ),
        None,
    )
    average_only = [
        metric
        for metric in metrics
        if "avg " in str(metric.get("name") or "").lower()
        and not (
            test_type.lower().startswith("strength:")
            and "range of motion" in str(metric.get("name") or "").lower()
        )
    ]
    ranked = sorted(
        enumerate(average_only),
        key=lambda item: (
            -metric_priority(test_type, item[1].get("name") or "", device),
            item[0],
        ),
    )
    rows: OrderedDict[tuple[str, str | None], dict] = OrderedDict()
    for _, metric in ranked:
        name = str(metric.get("name") or "")
        side_match = re.search(r"\s*\((Right|Left)\)\s*$", name, flags=re.IGNORECASE)
        base_name = name[: side_match.start()].strip() if side_match else name
        key = (display_metric_name(base_name), metric.get("unit"))
        if key not in rows:
            rows[key] = {
                "name": key[0],
                "value": metric.get("value") if not side_match else None,
                "unit": metric.get("unit"),
                "right_values": [],
                "left_values": [],
                "status": "neutral",
                "status_label": "Measured",
                "reference_note": "VALD Norms percentile not available through this API response.",
            }
        if side_match:
            side = side_match.group(1).lower()
            rows[key][f"{side}_values"].append(metric.get("value"))
            rows[key][f"{side}_value"] = average_numeric(rows[key][f"{side}_values"])
    selected = [
        row
        for row in list(rows.values())[:limit]
        if row.get("right_value") is not None and row.get("left_value") is not None
    ]
    for index, metric in enumerate(selected):
        asymmetry_value = bilateral_asymmetry(metric)
        if index == 0 and asymmetry is not None:
            asymmetry_value = abs(float(asymmetry["value"]))
        screened = add_metric_status({"name": "Asymmetry", "value": asymmetry_value})
        metric.update(
            {
                "asymmetry_value": asymmetry_value,
                "asymmetry_unit": "%",
                "direction": asymmetry_direction(metric),
                "status": screened.get("status"),
                "status_label": screened.get("status_label"),
                "reference_note": (
                    "VALD-reported asymmetry; operational <=10% screening band."
                    if index == 0 and asymmetry is not None
                    else "Calculated from displayed bilateral averages; operational <=10% screening band."
                ),
            }
        )
    return selected


def select_forcedecks_metrics(
    test_type: str, metrics: list[dict], limit: int = 5
) -> list[dict]:
    selected: list[dict] = []
    used: set[tuple[str, str | None]] = set()
    for display_name, patterns in forcedecks_metric_specs(test_type)[:limit]:
        candidates = matching_forcedecks_metrics(metrics, patterns)
        row = forcedecks_metric_row(display_name, candidates)
        if not row:
            continue
        key = (row["name"], row.get("unit"))
        if key in used:
            continue
        used.add(key)
        selected.append(row)
    return selected


def forcedecks_metric_specs(test_type: str) -> list[tuple[str, tuple[str, ...]]]:
    normalized = re.sub(r"[^a-z0-9]+", "", test_type.lower())
    if normalized in {"cmj", "abcmj"}:
        return [
            ("Jump Height (Imp-Mom)", ("jump height (imp-mom)",)),
            ("RSI-modified", ("rsi-modified",)),
            ("CMJ Stiffness", ("cmj stiffness", "lower-limb stiffness")),
            ("Concentric RFD - 200ms", ("concentric rfd - 200ms",)),
            ("Takeoff Peak Force", ("takeoff peak force",)),
        ]
    if normalized == "sj":
        return [
            ("Jump Height (Imp-Mom)", ("jump height (imp-mom)",)),
            ("RSI-modified", ("rsi-modified",)),
            ("Lower-Limb Stiffness", ("lower-limb stiffness", "landing stiffness")),
            ("Concentric RFD - 200ms", ("concentric rfd - 200ms",)),
            ("Takeoff Peak Force", ("takeoff peak force",)),
        ]
    if normalized == "imtp":
        return [
            ("Peak Force", ("peak force",)),
            ("Force at 100ms", ("force at 100ms",)),
            ("Force at 150ms", ("force at 150ms",)),
            ("Force at 200ms", ("force at 200ms",)),
            ("Baseline Force", ("baseline force",)),
            ("Peak RFD", ("peak rfd",)),
        ]
    if normalized == "dj":
        return [
            ("Jump Height (Imp-Mom)", ("jump height (imp-mom)",)),
            ("RSI", ("rsi",)),
            ("Contact Time", ("contact time",)),
            ("Takeoff Peak Force", ("takeoff peak force",)),
        ]
    if normalized in {"shldisoi", "shldisoy", "shldisot"}:
        return [
            ("Peak Vertical Force", ("peak vertical force",)),
            ("RFD - 200ms", ("rfd - 200ms",)),
            ("Start Time to Peak Force", ("start time to peak force",)),
        ]
    if normalized in {"slsb", "qsb"}:
        return [
            ("CoP Range ML", ("cop range ml", "cop range (ml)", "ml range")),
            ("CoP Range AP", ("cop range ap", "cop range (ap)", "ap range")),
            ("Mean Velocity", ("mean velocity",)),
            ("Total Excursion", ("total excursion",)),
            ("Area of CoP Ellipse", ("area of cop ellipse", "cop ellipse")),
        ]
    return [
        ("Jump Height (Imp-Mom)", ("jump height (imp-mom)",)),
        ("RSI-modified", ("rsi-modified",)),
        ("Concentric Mean Force", ("concentric mean force",)),
        ("Positive Takeoff Impulse", ("positive takeoff impulse",)),
        ("Takeoff Peak Force", ("takeoff peak force",)),
    ]


def matching_forcedecks_metrics(
    metrics: list[dict], patterns: tuple[str, ...]
) -> list[dict]:
    exact_matches: list[dict] = []
    contains_matches: list[dict] = []
    for metric in metrics:
        name = str(metric.get("name") or "")
        base = strip_side(name).lower()
        if should_skip_forcedecks_candidate(base):
            continue
        for pattern in patterns:
            pattern = pattern.lower()
            if base == pattern:
                exact_matches.append(metric)
                break
            if pattern in base:
                contains_matches.append(metric)
                break
    return exact_matches or contains_matches


def forcedecks_metric_row(display_name: str, candidates: list[dict]) -> dict | None:
    if not candidates:
        return None
    row = {
        "name": display_name,
        "unit": next((metric.get("unit") for metric in candidates if metric.get("unit")), None),
        "value_values": [],
        "right_values": [],
        "left_values": [],
        "status": "neutral",
        "status_label": "Measured",
        "reference_note": "Selected ForceDecks key metric; absolute quality needs VALD Hub norms or team benchmarks.",
    }
    for metric in candidates:
        name = str(metric.get("name") or "")
        side_match = re.search(r"\s*\((Right|Left)\)\s*$", name, flags=re.IGNORECASE)
        if side_match:
            side = side_match.group(1).lower()
            row[f"{side}_values"].append(metric.get("value"))
        else:
            row["value_values"].append(metric.get("value"))

    row["right_value"] = average_numeric(row["right_values"])
    row["left_value"] = average_numeric(row["left_values"])
    row["value"] = average_numeric(row["value_values"])
    if row["right_value"] is not None and row["left_value"] is not None:
        asymmetry_value = bilateral_asymmetry(row)
        screened = add_metric_status({"name": "Asymmetry", "value": asymmetry_value})
        row.update(
            {
                "asymmetry_value": asymmetry_value,
                "asymmetry_unit": "%",
                "direction": asymmetry_direction(row),
                "status": screened.get("status"),
                "status_label": screened.get("status_label"),
                "reference_note": "Calculated from displayed bilateral ForceDecks values; operational <=10% screening band.",
            }
        )
    elif row["value"] is None:
        return None
    return row


def strip_side(name: str) -> str:
    return re.sub(r"\s*\((Right|Left)\)\s*$", "", name, flags=re.IGNORECASE).strip()


def should_skip_forcedecks_candidate(base_name: str) -> bool:
    return any(
        term in base_name
        for term in (
            "ratio",
            " in inches",
            " / bm",
            " / bw",
            "relative",
            "flight time",
            "imp-dis",
            "50ms",
            "p1 ",
            "p2 ",
        )
    )


def metric_priority(test_type: str, name: str, device: str = "") -> int:
    label = name.lower()
    if "asymmetry" in label:
        return 120
    if device == "forcedecks":
        priorities = (
            ("concentric mean force", 125),
            ("eccentric mean force", 120),
            ("positive takeoff impulse", 115),
            ("concentric impulse", 110),
            ("landing rfd", 105),
            ("landing impulse", 100),
            ("stiffness", 95),
            ("mean landing force", 90),
            ("mean", 80),
        )
        return next((score for term, score in priorities if term in label), 20)
    if test_type.lower().startswith("strength:"):
        strength_priorities = (
            ("avg force", 115),
            ("avg rate of force development", 110),
            ("avg impulse", 105),
            ("avg time to peak force", 100),
            ("range of motion", 10),
        )
        return next((score for term, score in strength_priorities if term in label), 20)
    priorities = (
        ("avg range of motion", 110),
        ("avg force", 100),
        ("avg rate of force development", 90),
        ("avg impulse", 80),
        ("avg time to peak force", 70),
    )
    return next((score for term, score in priorities if term in label), 10)


def is_forcedecks_report_metric(name: str) -> bool:
    label = name.lower()
    if not re.search(r"\((right|left)\)\s*$", name, flags=re.IGNORECASE):
        return False
    if any(
        term in label
        for term in (
            "ratio",
            "peak",
            "maximum",
            "max ",
            " min ",
            "minimum",
            "50ms",
            "100ms",
            "p1 ",
            "p2 ",
        )
    ):
        return False
    if ":" in name:
        return False
    return any(
        term in label
        for term in (
            "concentric mean force",
            "eccentric mean force",
            "mean landing force",
            "positive takeoff impulse",
            "concentric impulse",
            "landing rfd",
            "landing impulse",
            "stiffness",
        )
    )


def average_numeric(values: list) -> float | None:
    numbers = []
    for value in values:
        try:
            numbers.append(float(value))
        except (TypeError, ValueError):
            continue
    if not numbers:
        return None
    return sum(numbers) / len(numbers)


def display_metric_name(name: str) -> str:
    return re.sub(r"\bAvg Range Of Motion\b", "Average ROM", name, flags=re.IGNORECASE).replace(
        "Avg ", "Average "
    )


def asymmetry_direction(metric: dict) -> str:
    right = metric.get("right_value")
    left = metric.get("left_value")
    if right is None or left is None:
        return "-"
    if float(right) == float(left):
        return "Equal"
    return "Towards Right" if float(right) > float(left) else "Towards Left"


def bilateral_asymmetry(metric: dict) -> float:
    right = abs(float(metric["right_value"]))
    left = abs(float(metric["left_value"]))
    larger = max(right, left)
    return round(abs(right - left) / larger * 100, 1) if larger else 0.0


def add_metric_status(metric: dict) -> dict:
    result = dict(metric)
    if "asymmetry" in str(metric.get("name") or "").lower():
        try:
            difference = abs(float(metric.get("value")))
        except (TypeError, ValueError):
            return result
        within_band = difference <= 10
        result.update(
            {
                "status": "green" if within_band else "red",
                "status_label": "Within screening band" if within_band else "Review asymmetry",
                "reference_note": "Operational asymmetry screening band: <=10%; not a VALD norm.",
            }
        )
    else:
        result.update(
            {
                "status": "neutral",
                "status_label": "Measured",
                "reference_note": "VALD Norms percentile not available through this API response.",
            }
        )
    return result


def format_metric_for_prompt(test_type: str, metric: dict) -> str:
    unit = metric.get("unit") or ""
    status = f" [{metric.get('status_label')}]" if metric.get("status_label") else ""
    if metric.get("right_value") is not None or metric.get("left_value") is not None:
        asymmetry = (
            f"; Asymmetry {metric.get('asymmetry_value')} {metric.get('asymmetry_unit')} "
            f"{metric.get('direction')}{status}"
            if metric.get("asymmetry_value") is not None
            else ""
        )
        return (
            f"{test_type}: {metric.get('name')} = Right {metric.get('right_value')} {unit}; "
            f"Left {metric.get('left_value')} {unit}{asymmetry}"
        ).strip()
    return (
        f"{test_type}: {metric.get('name')} = {metric.get('value')} {unit}{status}"
    ).strip()


def safe_doc_key(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_]+", "_", name).strip("_") or "Section"


def format_doc_value(value) -> str:
    if value is None:
        return "-"
    return format_metric_value_for_text(value)


def format_metric_value_for_text(value) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value or "-")
    return f"{number:,.2f}".rstrip("0").rstrip(".")


def compact_asymmetry_label(metric: dict) -> str:
    value = metric.get("asymmetry_value")
    if value is None:
        return "-"
    suffix = ""
    if metric.get("direction") == "Towards Right":
        suffix = "R"
    elif metric.get("direction") == "Towards Left":
        suffix = "L"
    return f"{format_metric_value_for_text(value)}{metric.get('asymmetry_unit') or '%'}{suffix}"


def add_markdown_to_doc(doc: Document, content: str) -> None:
    if not content.strip():
        doc.add_paragraph("")
        return
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading_match = re.match(r"^\*\*(.+?)\*\*$", line)
        if heading_match:
            doc.add_heading(heading_match.group(1), level=3)
            continue
        paragraph = doc.add_paragraph()
        parts = re.split(r"(\*\*.*?\*\*)", line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                paragraph.add_run(part[2:-2]).bold = True
            else:
                paragraph.add_run(part)


def paragraph_to_markdown(paragraph) -> str:
    text = ""
    for run in paragraph.runs:
        if not run.text:
            continue
        text += f"**{run.text}**" if run.bold else run.text
    if not text.strip():
        return ""
    style_name = (paragraph.style.name or "").lower()
    if "heading" in style_name:
        clean = text.replace("**", "").strip()
        if clean:
            return f"**{clean}**"
    return text.strip()


def extract_vald_interpretations(doc: Document, section_map: dict[str, str]) -> dict[str, str]:
    interpretations: dict[str, str] = {}
    current_key: str | None = None
    accumulator: list[str] = []

    def flush() -> None:
        nonlocal accumulator, current_key
        if current_key and current_key in section_map:
            content = "\n\n".join(line for line in accumulator if line.strip()).strip()
            interpretations[section_map[current_key]] = content
        accumulator = []
        current_key = None

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        start = re.search(r"\[START_VALD_INTERPRETATION_(.*?)\]", text)
        end = re.search(r"\[END_VALD_INTERPRETATION_(.*?)\]", text)
        if start:
            flush()
            current_key = start.group(1)
            remainder = text.replace(start.group(0), "").strip()
            if remainder:
                accumulator.append(remainder)
            continue
        if end:
            remainder = text.replace(end.group(0), "").strip()
            if remainder:
                accumulator.append(remainder)
            flush()
            continue
        if current_key:
            markdown = paragraph_to_markdown(paragraph)
            if markdown:
                accumulator.append(markdown)

    flush()
    return interpretations


def compact_interpretation(text: str, joint_name: str, sport: str) -> str:
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {
                "role": "system",
                "content": (
                    "You shorten technical sports reports without adding claims. "
                    "Follow the requested word limit exactly."
                ),
            },
            {
                "role": "user",
                "content": f"""
Condense this {joint_name} interpretation for a one-page {sport} report.
Use exactly these four Markdown headings:
**What Looks Good**
**Main Asymmetry**
**Why It Matters For {sport}**
**Performance Focus**

Maximum 180 words total. Retain one positive movement observation, the main asymmetry, why that joint action matters in {sport}, and how improving the flagged movement quality can support performance.
State briefly that numerical VALD Norms or approved benchmarks are needed for absolute interpretation.
Do not diagnose injury or introduce a new benchmark.

Original interpretation:
{text}
""",
            },
        ],
        temperature=0.1,
    )
    return constrain_interpretation_language(
        response.choices[0].message.content or text
    )


def joint_from_test_type(test_type: str) -> str:
    name = re.sub(r"^(ROM|Strength):\s*", "", test_type, flags=re.IGNORECASE)
    normalized = name.lower()
    for joint in (
        "Shoulder",
        "Elbow",
        "Wrist",
        "Hand",
        "Trunk",
        "Hip",
        "Knee",
        "Ankle",
        "Foot",
    ):
        if joint.lower() in normalized:
            return joint
    if any(term in normalized for term in ("jump", "squat", "countermovement")):
        return "Lower Body"
    if normalized in {"sj", "cmj", "abcmj", "dj", "hj"}:
        return "Lower Body"
    return "Whole Body"


def constrain_interpretation_language(content: str) -> str:
    replacements = (
        (r"\bclinically significant\b", "observed"),
        (r"\binsignificant\b", "observed"),
        (r"\bsignificant\b", "observed"),
        (r"\bsignificantly\b", "observably"),
        (r"\babnormal\b", "observed"),
        (r"\bdeficien(?:t|cy)\b", "difference"),
        (r"\bdeficit\b", "difference"),
        (r"\bdeficits\b", "differences"),
        (r"\bcommendable\b", "good from a symmetry perspective"),
        (r"\bsolid explosive power\b", "useful jump output"),
        (r"\bsolid explosive capability\b", "reported explosive output"),
        (r"\bsolid foundation in vertical power\b", "reported vertical jump output"),
        (r"\bgood from a symmetry perspective symmetry\b", "well-balanced symmetry"),
        (r"\bjump height .*? are good from a symmetry perspective\b", "jump height values are reported"),
        (r"\bexcellent\b", "good from a symmetry perspective"),
        (r"\bacceptable limits\b", "the operational screening band"),
        (
            r"\b(?:elevated |increased |potential )?risk of (?:an? )?(?:overuse )?injur(?:y|ies)\b",
            "consideration for qualified practitioner review",
        ),
        (r"\binjuries\b", "clinical concerns"),
        (r"\binjury\b", "clinical concern"),
        (r"\bpotential areas\b", "areas"),
        (r"\bpotential area\b", "area"),
        (r"\bpotential imbalances\b", "observed asymmetry patterns"),
        (r"\bpotential imbalance\b", "observed asymmetry pattern"),
        (r"\ba area\b", "an area"),
        (r"\breturn-to-play\b", "sport participation"),
        (r"\b\d+\s*-\s*\d+\s*weeks?\b", "a planned retest window"),
        (r"\btargeted interventions\b", "practitioner-led training decisions"),
        (r"\binterventions\b", "training decisions"),
        (r"\bis recommended\b", "can be considered by the practitioner"),
        (r"\bare recommended\b", "can be considered by the practitioner"),
        (r"\bprevent compensatory patterns\b", "observe compensatory patterns"),
        (r"\bprevent\b", "observe"),
        (r"\breduce potential compensatory patterns\b", "observe compensatory patterns"),
        (r"\breduce compensatory patterns\b", "observe compensatory patterns"),
        (r"\bFurther assessment should occur\b", "Further assessment can be considered"),
        (r"\bTo high-quality performance\b", "For performance"),
        (r"\bto high-quality performance\b", "for performance"),
        (
            r"\bit can be considered by the practitioner that ([^.]+?) focuses on\b",
            r"\1 can focus on",
        ),
    )
    safe_content = content
    for pattern, replacement in replacements:
        safe_content = re.sub(pattern, replacement, safe_content, flags=re.IGNORECASE)
    safe_content = re.sub(r"\ba observed\b", "an observed", safe_content, flags=re.IGNORECASE)
    return re.sub(
        r"(?:may |could |can )?(?:increase|elevate|heighten) (?:the )?consideration for qualified practitioner review",
        "warrants qualified practitioner review",
        safe_content,
        flags=re.IGNORECASE,
    )
