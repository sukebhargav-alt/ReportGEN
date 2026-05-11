from fastapi import APIRouter, UploadFile, File, Body, Form
from fastapi.responses import StreamingResponse
from io import BytesIO
import json
import re
from docx import Document
from modules.common.config import client
from modules.profiling.parser import parse_full_profiling_file
from modules.profiling.renderer import render_premium_pdf
from modules.profiling.prompts import build_profiling_section_prompt

router = APIRouter()

# Sheet names excluded from all reporting pipelines
SKIPPED_SECTIONS = {"vald raw data"}

def _is_skipped(section_name: str) -> bool:
    return section_name.strip().lower() in SKIPPED_SECTIONS

@router.post("/process_profiling_excel")
async def process_profiling_excel(file: UploadFile = File(...)):
    contents = await file.read()
    excel_data = BytesIO(contents)
    parsed = parse_full_profiling_file(excel_data)
    return parsed

@router.post("/generate_profiling_interpretations")
async def generate_profiling_interpretations(data: dict = Body(...)):
    profile = data.get("profile", {})
    assessments = data.get("assessments", {})
    interpretations = {}

    for section_name, entries in assessments.items():
        if not entries or _is_skipped(section_name): continue
        formatted_entries = []
        last_group = None

        for entry in entries:
            label = entry.get("Test Name") or entry.get("Metric Name")
            unit = entry.get("Unit", "") or ""
            test_group = entry.get("Test Group", "") or ""

            # Emit a sub-group header whenever the Test Group changes
            if test_group and test_group != last_group:
                formatted_entries.append(f"\n=== {test_group} ===")
                last_group = test_group

            if entry.get("Left") is not None or entry.get("Right") is not None:
                right = entry.get("Right")
                left  = entry.get("Left")
                asym  = entry.get("Asymmetry")
                # Only include sides that actually have data
                parts = []
                if right is not None: parts.append(f"R {right}")
                if left  is not None: parts.append(f"L {left}")
                asym_str = f" / Asym {asym}" if asym else ""
                sides_str = " / ".join(parts)
                formatted_entries.append(f"  {label}: {sides_str} {unit}{asym_str}")
            else:
                value = entry.get("Value")
                formatted_entries.append(f"  {label}: {value} {unit}")

        formatted_entries_text = "\n".join(formatted_entries)
        section_prompt = build_profiling_section_prompt(section_name, profile, formatted_entries_text)
        
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": "You are an elite sports performance scientist writing high-level technical performance reports."},
                {"role": "user", "content": section_prompt}
            ],
            temperature=0.25
        )
        interpretations[section_name] = response.choices[0].message.content
    return interpretations

@router.post("/generate_profiling_word_template")
async def generate_profiling_word_template(data: dict = Body(...)):
    profile = data.get("profile", {})
    assessments = data.get("assessments", {})
    interpretations = data.get("interpretations", {})
    doc = Document()

    # ── Profile section ───────────────────────────────────────────────
    doc.add_heading("Athlete Performance Report", level=1)
    doc.add_heading("Athlete Profile (DO NOT EDIT)", level=2)
    for key, value in profile.items():
        if value not in [None, "", "None"]:
            doc.add_paragraph(f"{key}: {value}")

    # ── Assessment sections ───────────────────────────────────────────
    for section_name, entries in assessments.items():
        if not entries or _is_skipped(section_name): continue
        doc.add_page_break()
        doc.add_heading(section_name, level=2)

        # Determine which columns are actually used in this section
        has_bilateral = any(
            e.get("Left") is not None or e.get("Right") is not None
            for e in entries
        )
        has_scalar = any(
            e.get("Left") is None and e.get("Right") is None and e.get("Value") is not None
            for e in entries
        )
        has_asym = any(e.get("Asymmetry") is not None for e in entries)

        # Build column header list
        col_headers = ["Test Group", "Test Name"]
        if has_bilateral:
            col_headers += ["Left", "Right"]
        if has_asym:
            col_headers.append("Asymmetry")
        if has_scalar:
            col_headers.append("Value")
        col_headers.append("Unit")

        # Create Word table
        table = doc.add_table(rows=1, cols=len(col_headers))
        table.style = "Table Grid"

        # Header row
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(col_headers):
            hdr_cells[i].text = h

        # Data rows
        for entry in entries:
            label = entry.get("Test Name") or entry.get("Metric Name") or ""
            unit = str(entry.get("Unit", "") or "")
            test_group = str(entry.get("Test Group", "") or "")
            row_cells = table.add_row().cells
            col_idx = 0
            row_cells[col_idx].text = test_group; col_idx += 1
            row_cells[col_idx].text = label; col_idx += 1
            if has_bilateral:
                row_cells[col_idx].text = str(entry.get("Left") or "-"); col_idx += 1
                row_cells[col_idx].text = str(entry.get("Right") or "-"); col_idx += 1
            if has_asym:
                row_cells[col_idx].text = str(entry.get("Asymmetry") or "-"); col_idx += 1
            if has_scalar:
                row_cells[col_idx].text = str(entry.get("Value") or "-"); col_idx += 1
            row_cells[col_idx].text = unit; col_idx += 1

        doc.add_paragraph()  # spacer

        # Interpretation block with sentinels
        safe_section = re.sub(r"[^a-zA-Z0-9_]", "_", section_name)
        doc.add_paragraph(f"[START_INTERPRETATION_{safe_section}]")
        doc.add_paragraph(interpretations.get(section_name, ""))
        doc.add_paragraph(f"[END_INTERPRETATION_{safe_section}]")

    # ── Final Recommendations ─────────────────────────────────────────
    doc.add_page_break()
    doc.add_heading("Final Recommendations", level=2)
    doc.add_paragraph("[START_FINAL_RECOMMENDATIONS]")
    doc.add_paragraph("[END_FINAL_RECOMMENDATIONS]")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=Athlete_Report_Draft.docx"}
    )

def _get_paragraph_markdown(p):
    md = ""
    for run in p.runs:
        text = run.text
        if not text:
            continue
        if run.bold:
            text = f"**{text}**"
        if run.italic:
            text = f"_{text}_"
        md += text
    
    style = p.style.name.lower()
    text_stripped = p.text.strip()
    
    # Aggressive list detection
    is_list = "bullet" in style or "number" in style or "list" in style
    # Also detect if it looks like a numbered list manually typed
    is_manual_number = re.match(r"^\d+[\.\)]\s+", text_stripped)
    
    if is_list:
        if "bullet" in style:
            if not md.strip().startswith("*"):
                md = f"* {md.lstrip()}"
        else:
            # Assume numbered for "number" or "list" styles
            if not re.match(r"^\d+\.\s+", md.strip()):
                md = f"1. {md.lstrip()}"
    elif is_manual_number:
        # Ensure it's treated as a markdown list by having exactly one number dot space
        content = re.sub(r"^\d+[\.\)]\s+", "", md).lstrip()
        md = f"1. {content}"
        
    if "heading" in style:
        level_match = re.search(r"\d", style)
        level = int(level_match.group()) if level_match else 1
        md = f"{'#' * level} {md}"
        
    return md

@router.post("/parse_profiling_edited_word_file")
async def parse_profiling_edited_word_file(file: UploadFile = File(...), original_data: str = Form(...)):
    contents = await file.read()
    doc = Document(BytesIO(contents))
    
    original_parsed = json.loads(original_data)
    assessments = original_parsed.get("assessments", {})
    
    interpretations = {}
    final_recommendations = ""
    
    # Pre-calculate safe section names for matching
    section_map = {re.sub(r"[^a-zA-Z0-9_]", "_", name): name for name in assessments.keys()}
    
    current_section = None
    collecting_interpretation = False
    collecting_recommendations = False
    
    interp_accumulator = []
    recs_accumulator = []

    for p in doc.paragraphs:
        md_line = _get_paragraph_markdown(p)
        text = p.text.strip()
        
        # ── Final Recommendations ─────────────────────────────────────
        if "[START_FINAL_RECOMMENDATIONS]" in text:
            collecting_recommendations = True
            md_line = md_line.replace("[START_FINAL_RECOMMENDATIONS]", "")
        
        if "[END_FINAL_RECOMMENDATIONS]" in text:
            md_line = md_line.replace("[END_FINAL_RECOMMENDATIONS]", "")
            if collecting_recommendations:
                if md_line.strip(): recs_accumulator.append(md_line)
                final_recommendations = "\n".join(recs_accumulator).strip()
            collecting_recommendations = False
            continue

        # ── Interpretations ───────────────────────────────────────────
        re_interp_start = re.search(r"\[START_INTERPRETATION_(.*?)\]", text)
        if re_interp_start:
            safe_name = re_interp_start.group(1)
            if safe_name in section_map:
                current_section = section_map[safe_name]
                collecting_interpretation = True
                interp_accumulator = []
                # Remove the sentinel from the markdown line
                md_line = md_line.replace(re_interp_start.group(0), "")

        if current_section:
            end_sentinel = f"[END_INTERPRETATION_{re.sub(r'[^a-zA-Z0-9_]', '_', current_section)}]"
            if end_sentinel in text:
                md_line = md_line.replace(end_sentinel, "")
                if md_line.strip(): interp_accumulator.append(md_line)
                interpretations[current_section] = "\n".join(interp_accumulator).strip()
                collecting_interpretation = False
                current_section = None
                continue

        # ── Accumulation ──────────────────────────────────────────────
        if collecting_recommendations:
            line = md_line.strip()
            if line:
                # Ensure every line in recommendations is treated as a list item if it's not already
                if not line.startswith(('*', '1.')):
                    line = f"* {line}"
                recs_accumulator.append(line)
        elif collecting_interpretation:
            if md_line.strip(): interp_accumulator.append(md_line)

    final_recommendations = "\n".join(recs_accumulator).strip()
    return {"interpretations": interpretations, "final_recommendations": final_recommendations}

@router.post("/generate_profiling_final_pdf")
async def generate_profiling_final_pdf(data: dict = Body(...)):
    buffer = BytesIO()
    await render_premium_pdf(buffer, data)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=Premium_Athlete_Report.pdf"})
