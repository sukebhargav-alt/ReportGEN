from fastapi import APIRouter, UploadFile, File, Body, Form, HTTPException
from fastapi.responses import StreamingResponse
from io import BytesIO
import asyncio
import json
import logging
import re
from docx import Document
from modules.common.config import client
from modules.cpet.parser import parse_full_cpet_file
from modules.cpet.renderer import render_cpet_premium_pdf
from modules.cpet.prompts import build_cpet_interpretation_prompt
from modules.cpet.comparison_prompts import build_comparison_prompt
from modules.cpet.comparison_renderer import render_comparison_pdf

router = APIRouter()
logger = logging.getLogger(__name__)

@router.post("/process_cpet_excel")
async def process_cpet_excel(file: UploadFile = File(...)):
    contents = await file.read()
    excel_data = BytesIO(contents)
    parsed = parse_full_cpet_file(excel_data)
    return parsed

@router.post("/generate_cpet_interpretations")
async def generate_cpet_interpretations(data: dict = Body(...)):
    profile = data.get("profile", {})
    results = data.get("results", [])

    if not isinstance(results, list) or not results:
        raise HTTPException(
            status_code=400,
            detail=(
                "No CPET result data was supplied. Process a CPET file before "
                "generating ACSM insights."
            ),
        )

    systems = [
        "Overall Interpretation",
        "Ventilation System",
        "Cardiovascular System",
        "Ventilatory Perfusion",
        "Metabolic System",
    ]

    semaphore = asyncio.Semaphore(2)

    async def fetch_system(system_name: str) -> tuple[str, str, str | None]:
        async with semaphore:
            try:
                response = await asyncio.to_thread(
                    client.chat.completions.create,
                    model="gpt-4o-mini",
                    messages=[
                        {
                            "role": "system",
                            "content": "You are an elite sports performance scientist leveraging ACSM standards.",
                        },
                        {"role": "user", "content": build_cpet_interpretation_prompt(system_name, profile, results)},
                    ],
                    temperature=0.25,
                )
                return system_name, response.choices[0].message.content or "", None
            except Exception as exc:
                logger.exception("Failed to generate CPET ACSM insight for %s", system_name)
                return system_name, "", str(exc)

    results_gathered = await asyncio.gather(*[fetch_system(s) for s in systems])
    interpretations: dict[str, str] = {}
    warnings: list[dict[str, str]] = []

    for system_name, text, error in results_gathered:
        if error:
            warnings.append(
                {
                    "section": system_name,
                    "message": "AI service error while generating this section.",
                }
            )
            interpretations[system_name] = (
                f"{system_name} insight could not be generated on this attempt. "
                "Please retry after confirming the backend AI service is available."
            )
        else:
            interpretations[system_name] = text

    if len(warnings) == len(systems):
        raise HTTPException(
            status_code=502,
            detail=(
                "Unable to generate ACSM insights because the backend AI service "
                "failed for every section. Check OPENAI_API_KEY and Render backend logs, then retry."
            ),
        )

    return {"insights": interpretations, "warnings": warnings}

@router.post("/generate_cpet_final_pdf")
async def generate_cpet_final_pdf(data: dict = Body(...)):
    buffer = BytesIO()
    await render_cpet_premium_pdf(buffer, data)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=Athlete_CPET_Final_Report.pdf"}
    )

@router.post("/generate_cpet_word_template")
async def generate_cpet_word_template(payload: dict = Body(...)):
    import base64
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    data        = payload.get("data", {})
    graphs_b64  = payload.get("graphsBase64", [])   # list of 9 base64 strings
    insights    = payload.get("interpretations", {}) # dict keyed by section name

    profile = data.get("profile", {})
    results = data.get("results", [])

    doc = Document()
    doc.add_heading("Cardiopulmonary Exercise Testing (CPET) Report", level=1)
    doc.add_heading("Athlete Details", level=2)

    # Build BMI
    bmi = profile.get("BMI (kg/m2)")
    if not bmi:
        try:
            w = float(profile.get("Weight (kg)", 0) or 0)
            h = float(profile.get("Height (cm)", 0) or 0) / 100
            if h > 0:
                bmi = round(w / (h * h), 2)
        except Exception:
            bmi = None

    display_fields = [
        ("First Name",        profile.get("First Name", "-")),
        ("Last Name",         profile.get("Last Name",  "-")),
        ("Gender",            profile.get("Gender",     "-")),
        ("Age",               str(round(float(profile["Age"]))) if profile.get("Age") else "-"),
        ("Height",            f"{profile['Height (cm)']} cm" if profile.get("Height (cm)") else "-"),
        ("Weight",            f"{profile['Weight (kg)']} kg" if profile.get("Weight (kg)") else "-"),
        ("BMI",               str(bmi) if bmi else "-"),
        ("Protocol & HR Max", f"Protocol: {profile.get('Protocol', '-')} | HR Max: {profile.get('HR Max', '-')}"),
        ("Exercise Duration", profile.get("Exercise Duration", "-")),
        ("Sport",             profile.get("Sport", "-")),
        ("Injuries",          profile.get("Injuries") or "-"),
    ]

    profile_table = doc.add_table(rows=0, cols=2)
    profile_table.style = "Table Grid"
    for label, value in display_fields:
        row = profile_table.add_row().cells
        run0 = row[0].paragraphs[0].add_run(label)
        run0.bold = True
        row[1].text = str(value)

    def add_markdown_insight(doc_obj, text, base_level=3):
        import re
        for line in text.split('\n'):
            line = line.strip()
            if not line:
                continue
            if line.startswith("### "):
                doc_obj.add_heading(line[4:].replace('**', ''), level=base_level)
            elif line.startswith("## "):
                doc_obj.add_heading(line[3:].replace('**', ''), level=max(1, base_level-1))
            elif line.startswith("# "):
                doc_obj.add_heading(line[2:].replace('**', ''), level=max(1, base_level-2))
            else:
                p = doc_obj.add_paragraph()
                parts = re.split(r'(\*\*.*?\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)

    # The Overall Interpretation is handled below in the dedicated section

    SECTIONS = [
        ("Ventilation System", [
            (3, "VE vs Time"),
            (4, "VE vs VCO2"),
            (5, "VE/VO2 & VE/VCO2 vs Time"),
        ]),
        ("Cardiovascular System", [
            (0, "HR vs VO2"),
            (2, "VO2 & VCO2 vs Time"),
        ]),
        ("Ventilatory Perfusion", [
            (6, "VT vs VE"),
            (7, "RQ vs Time"),
            (8, "PetO2 / PetCO2 / PaCO2_e vs Time"),
        ]),
        ("Metabolic System", [
            (7, "RQ vs Time"),
            (1, "RQ & HR vs Time"),
        ]),
    ]

    SYSTEM_METRIC_MAP = {
        "Ventilation System": ["VE", "VO2", "VO2/kg", "VCO2", "RR", "RF", "VT", "O2 pulse"],
        "Cardiovascular System": ["HR", "HRR"],
        "Ventilatory Perfusion": ["VT1", "VT2", "PetO2", "PetCO2"],
        "Metabolic System": ["RQ", "EEh", "METS", "RER"],
    }

    # Build Overall Interpretation Section
    doc.add_heading("Overall Interpretation", level=2)
    doc.add_paragraph("[START_INSIGHT_Overall_Interpretation]")
    if insights.get("Overall Interpretation"):
        add_markdown_insight(doc, insights["Overall Interpretation"], base_level=3)
    
    # Add VO2 Max section if available
    if data.get("vo2_max_gauge"):
        vg = data["vo2_max_gauge"]
        doc.add_heading(f"Aerobic Capacity ({vg.get('label', 'VO2 Max')})", level=3)
        doc.add_paragraph(f"Measured Value: {vg.get('value')} {vg.get('unit', '')}")
        doc.add_paragraph(f"Performance Category: {vg.get('category')}")
        doc.add_paragraph(f"Reference Range: {vg.get('range_label')}")
    
    doc.add_paragraph("[END_INSIGHT_Overall_Interpretation]")
    doc.add_page_break()

    for section_title, chart_slots in SECTIONS:
        section_insight = insights.get(section_title, "") if isinstance(insights, dict) else ""
        rq_hr_insight = ""
        
        # Split insights if needed
        aerobic_insight = ""
        vt_thresholds_insight = ""
        rq_hr_insight = ""

        if section_title == "Cardiovascular System" and section_insight and "Aerobic Capacity Interpretation" in section_insight:
            parts = re.split(r'(###\s*Aerobic\s*Capacity\s*Interpretation|##\s*Aerobic\s*Capacity\s*Interpretation|#\s*Aerobic\s*Capacity\s*Interpretation)', section_insight, flags=re.IGNORECASE)
            if len(parts) >= 3:
                section_insight = parts[0].strip()
                aerobic_insight = (parts[1] + parts[2]).strip()

        if section_title == "Ventilation System" and section_insight and "Ventilatory Thresholds Interpretation" in section_insight:
            parts = re.split(r'(###\s*Ventilatory\s*Thresholds\s*Interpretation|##\s*Ventilatory\s*Thresholds\s*Interpretation|#\s*Ventilatory\s*Thresholds\s*Interpretation)', section_insight, flags=re.IGNORECASE)
            if len(parts) >= 3:
                section_insight = parts[0].strip()
                vt_thresholds_insight = (parts[1] + parts[2]).strip()

        if section_title == "Metabolic System" and section_insight and "RQ & HR vs Time Interpretation" in section_insight:
            parts = re.split(r'(###\s*RQ\s*&\s*HR\s*vs\s*Time\s*Interpretation|##\s*RQ\s*&\s*HR\s*vs\s*Time\s*Interpretation|#\s*RQ\s*&\s*HR\s*vs\s*Time\s*Interpretation)', section_insight, flags=re.IGNORECASE)
            if len(parts) >= 3:
                section_insight = parts[0].strip()
                rq_hr_insight = (parts[1] + parts[2]).strip()

        doc.add_heading(section_title, level=2)
        if results:
            relevant_keys = SYSTEM_METRIC_MAP.get(section_title, [])
            section_metrics = []
            for row in results:
                param = str(row.get("Parameter", "")).strip()
                if not param: continue
                matched = False
                for k in relevant_keys:
                    if k.lower() in param.lower():
                        matched = True
                        break
                if matched:
                    section_metrics.append(row)
            
            if section_metrics:
                doc.add_heading("Key Metrics", level=3)
                param_table = doc.add_table(rows=0, cols=2)
                param_table.style = "Table Grid"
                for sm in section_metrics:
                    r = param_table.add_row().cells
                    param_name = str(sm.get("Parameter", ""))
                    unit = ""
                    val_str = []
                    import math
                    for k, v in sm.items():
                        if k == "Parameter": continue
                        if isinstance(v, float) and math.isnan(v): continue
                        safe_v = str(v).strip()
                        if safe_v in ["", "nan", "None", "---", "undefined"]: continue
                        if k.lower() == "um": unit = safe_v
                        else: val_str.append(f"{k}: {safe_v}")
                    if unit: param_name = f"{param_name} ({unit})"
                    r[0].paragraphs[0].add_run(param_name).bold = True
                    r[1].text = " | ".join(val_str)
                doc.add_paragraph()

        for idx, chart_title in chart_slots:
            if idx < len(graphs_b64) and graphs_b64[idx]:
                try:
                    img_bytes  = base64.b64decode(graphs_b64[idx])
                    img_buffer = BytesIO(img_bytes)
                    doc.add_picture(img_buffer, width=Inches(6.5), height=Inches(2.0))
                    cap = doc.add_paragraph(chart_title)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.runs[0].italic = True
                    cap.runs[0].font.size = Pt(9)
                    cap.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

                    if section_title == "Cardiovascular System" and chart_title == "VO2 & VCO2 vs Time" and aerobic_insight:
                        doc.add_paragraph("[START_AEROBIC_CAPACITY_INSIGHT]")
                        add_markdown_insight(doc, aerobic_insight, base_level=4)
                        doc.add_paragraph("[END_AEROBIC_CAPACITY_INSIGHT]")

                    if section_title == "Ventilation System" and chart_title == "VE/VO2 & VE/VCO2 vs Time" and vt_thresholds_insight:
                        doc.add_paragraph("[START_VENTILATORY_THRESHOLDS_INSIGHT]")
                        add_markdown_insight(doc, vt_thresholds_insight, base_level=4)
                        doc.add_paragraph("[END_VENTILATORY_THRESHOLDS_INSIGHT]")

                    if section_title == "Metabolic System" and chart_title == "RQ & HR vs Time":
                        if rq_hr_insight:
                            doc.add_paragraph("[START_RQ_HR_INSIGHT]")
                            add_markdown_insight(doc, rq_hr_insight, base_level=4)
                            doc.add_paragraph("[END_RQ_HR_INSIGHT]")
                        
                        # Add a 4x4 empty table below the graph
                        doc.add_paragraph()
                        table = doc.add_table(rows=4, cols=4)
                        table.style = "Table Grid"
                        
                        # Populate headers
                        hdr_cells = table.rows[0].cells
                        hdr_cells[0].text = "Sr No"
                        hdr_cells[1].text = "Range For RQ"
                        hdr_cells[2].text = "Hr"
                        hdr_cells[3].text = "Metabolism"
                        
                        # Make headers bold
                        for cell in hdr_cells:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True
                        
                        doc.add_paragraph()
                        doc.add_paragraph()

                except Exception as e:
                    print(f"[CPET Word] Failed to embed chart {idx} ({chart_title}): {e}")
            else:
                doc.add_paragraph(f"[Chart not available: {chart_title}]")

        safe_title = re.sub(r"[^a-zA-Z0-9_]", "_", section_title)
        doc.add_paragraph()
        doc.add_heading(f"{section_title} — ACSM Insights", level=3)
        doc.add_paragraph(f"[START_INSIGHT_{safe_title}]")
        if section_insight:
            add_markdown_insight(doc, section_insight, base_level=4)
        doc.add_paragraph("[END_INSIGHT_" + safe_title + "]")

    doc.add_page_break()
    doc.add_heading("Final Recommendations", level=1)
    doc.add_paragraph()

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=Athlete_CPET_Report_Draft.docx"},
    )
    
@router.post("/extract_final_recommendations")
async def extract_final_recommendations(file: UploadFile = File(...)):
    if not file.filename or not file.filename.endswith(".docx"):
        return {"error": "Invalid file format. Please upload a .docx file."}
    
    content = await file.read()
    doc = Document(BytesIO(content))

    # ── Extraction Logic ──────────────────────────────────────────────────────
    full_text = "\n".join([p.text for p in doc.paragraphs])
    
    sections_map = {
        "Overall_Interpretation": "Overall Interpretation",
        "Ventilation_System": "Ventilation System",
        "Cardiovascular_System": "Cardiovascular System",
        "Ventilatory_Perfusion": "Ventilatory Perfusion",
        "Metabolic_System": "Metabolic System",
    }
    
    # 1. Extract Main Insights
    insights: dict[str, str] = {}
    for safe_key, display_name in sections_map.items():
        pattern = rf"\[\s*START_INSIGHT_{safe_key}\s*\](.*?)\[\s*END_INSIGHT_{safe_key}\s*\]"
        match = re.search(pattern, full_text, re.DOTALL | re.IGNORECASE)
        content = match.group(1).strip() if match else ""
        
        # Clean out any sub-insight tags that might be nested inside
        content = re.sub(r"\[\s*START_.*?_INSIGHT\s*\].*?\[\s*END_.*?_INSIGHT\s*\]", "", content, flags=re.DOTALL | re.IGNORECASE)
        
        # CLEANUP: Remove stray numbers (1., 2., 3.) that often appear before/after headers in Word
        # First, strip any leading numbers from the very start of the content
        content = re.sub(r"^\d+[\.\)]\s*", "", content).strip()
        
        for h in ["Data Interpretation", "Performance Risk Assessment", "Key Findings"]:
            # This regex finds "1. Data Interpretation", "Data Interpretation:", etc. and turns it into a bullet
            pattern_h = rf"(?:\d+[\.\)]\s*)?{re.escape(h)}[:\s-]*"
            content = re.sub(rf"(?i){pattern_h}", f"\n\n* **{h}**: ", content)
            
        insights[display_name] = content.strip()

    # 2. Extract Sub-Insights
    sub_map = {
        "AEROBIC_CAPACITY": ("Cardiovascular System", "Aerobic Capacity Interpretation"),
        "VENTILATORY_THRESHOLDS": ("Ventilation System", "Ventilatory Thresholds Interpretation"),
        "RQ_HR": ("Metabolic System", "RQ & HR vs Time Interpretation")
    }
    
    for tag_key, (parent_section, header_name) in sub_map.items():
        pattern = rf"\[\s*START_{tag_key}_INSIGHT\s*\](.*?)\[\s*END_{tag_key}_INSIGHT\s*\]"
        match = re.search(pattern, full_text, re.DOTALL | re.IGNORECASE)
        if match and match.group(1).strip():
            content = match.group(1).strip()
            # Remove leading numbers from sub-insights
            content = re.sub(r"^\d+[\.\)]\s*", "", content).strip()
            
            # DEDUPLICATE: If the content already starts with the header name, strip it
            content = re.sub(rf"^{re.escape(header_name)}[:\s-]*", "", content, flags=re.IGNORECASE).strip()
            
            # Ensure header is on its own line at the start of the text for the renderer
            if insights.get(parent_section):
                insights[parent_section] += f"\n\n### {header_name}\n{content}"
            else:
                insights[parent_section] = f"### {header_name}\n{content}"

    # 3. Extract Final Recommendations
    recommendations_text = []
    capture = False
    for para in doc.paragraphs:
        # Strip tabs and leading/trailing spaces
        t = para.text.replace("\t", " ").strip()
        if "Final Recommendations" in t and not capture:
            capture = True
            continue
        if capture and t:
            if t.startswith("[") and t.endswith("]"): continue
            # Format as markdown bullet points
            recommendations_text.append(f"* {t}")

    # 4. Extract the Metabolic Table separately
    rq_table_html = ""
    for table in doc.tables:
        if len(table.columns) == 4 and len(table.rows) >= 3:
            # Check if the table is effectively empty (excluding header)
            is_empty = True
            for r_idx in range(1, len(table.rows)):
                for cell in table.rows[r_idx].cells:
                    if cell.text.strip():
                        is_empty = False
                        break
                if not is_empty: break
            
            if is_empty: continue

            html = '<table style="width:100%; border-collapse: collapse; margin-top: 24px; border: 1px solid #e5e7eb;">'
            for r_idx, row in enumerate(table.rows):
                bg = "background: #f8fafc; font-weight: bold;" if r_idx == 0 else ""
                html += f'<tr style="{bg}">'
                for cell in row.cells:
                    html += f'<td style="border: 1px solid #e5e7eb; padding: 10px; font-size: 11px; color: #374151;">{cell.text.strip()}</td>'
                html += "</tr>"
            html += "</table>"
            rq_table_html = html
            break

    # Fallback keys for the template
    if "Metabolic System" in insights:
        insights["Metabolic Efficiency (RQ vs HR)"] = insights["Metabolic System"]

    return {
        "recommendations": "\n".join(recommendations_text).strip(),
        "insights": insights,
        "rq_table": rq_table_html
    }


# ── Comparison Interpretations ────────────────────────────────────────────────

@router.post("/generate_comparison_interpretations")
async def generate_comparison_interpretations(data: dict = Body(...)):
    """
    Generate AI-powered What/Why/How interpretations for each physiological
    category in a CPET comparison.

    Payload:
      profile  – athlete profile dict
      metrics  – list of ComparisonMetric-like dicts
                 (label, category, baseline, followUp, delta, deltaPercent,
                  improvement, unit)
    """
    profile = data.get("profile", {})
    metrics = data.get("metrics", [])

    CATEGORIES = ["overall", "ventilation", "cardiovascular", "perfusion", "metabolic"]

    async def fetch_category(category: str) -> tuple[str, dict]:
        if category == "overall":
            cat_metrics = [
                m for m in metrics
                if m.get("baseline") is not None
                and m.get("followUp") is not None
            ]
        else:
            cat_metrics = [
                m for m in metrics
                if m.get("category") == category
                and m.get("baseline") is not None
                and m.get("followUp") is not None
            ]

        if not cat_metrics:
            return category, {"summary": "No data available for this system.", "metrics": []}

        prompt = build_comparison_prompt(category, profile, cat_metrics)
        response = await asyncio.to_thread(
            client.chat.completions.create,
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are an elite sports physiologist. "
                        "Return ONLY the requested JSON — no markdown, no prose outside the JSON."
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            temperature=0.20,
            response_format={"type": "json_object"},
        )
        raw = response.choices[0].message.content or "{}"
        try:
            parsed = json.loads(raw)
        except Exception:
            parsed = {"summary": raw, "metrics": []}
        return category, parsed

    results = await asyncio.gather(*[fetch_category(c) for c in CATEGORIES])
    return {"insights": {cat: insight for cat, insight in results}}


@router.post("/generate_comparison_pdf")
async def generate_comparison_pdf(data: dict = Body(...)):
    """
    Generate the premium 6-8 page CPET progress comparison PDF.

    Payload mirrors the comparison dashboard state:
      profile, metrics, insights, overall_score, improved_count,
      declined_count, unchanged_count, baseline_date, followup_date,
      radar_b64 (optional), final_recommendations (optional)
    """
    buffer = BytesIO()
    await render_comparison_pdf(buffer, data)
    buffer.seek(0)
    athlete = data.get("profile", {}).get(
        "Name",
        data.get("profile", {}).get("Athlete Name", "Athlete")
    )
    safe_name = re.sub(r"[^a-zA-Z0-9_-]", "_", athlete)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="{safe_name}_Progress_Report.pdf"'
        },
    )
